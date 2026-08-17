"""Bank-Onto-Bench v2 → 엑셀 정리본 (HF 데이터셋 카드 스타일).

실행: python bench/export_xlsx.py
출력: docs/bank-onto-bench-v2.xlsx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

REPO_ROOT = Path(__file__).resolve().parent.parent
DATASET = REPO_ROOT / "data" / "bank-onto-bench-v2.jsonl"
OUTPUT = REPO_ROOT / "docs" / "bank-onto-bench-v2.xlsx"

HEADER_FILL = PatternFill("solid", start_color="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(vertical="top", wrap_text=True)


def style(ws, widths):
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "A2"


def main() -> int:
    rows = [json.loads(l) for l in DATASET.read_text().splitlines() if l]
    wb = Workbook()

    ws = wb.active
    ws.title = "dataset_card"
    card = [
        ("항목", "내용"),
        ("dataset_name", "Bank-Onto-Bench"),
        ("version", "2.0"),
        ("size", "1,000문항 (공시형 700 : 규정형 300 = 7:3)"),
        ("description",
         "한국 은행 도메인 QA 벤치마크. 정답은 전부 공개 원문(금감원 공시 필드, "
         "법령 조문, 시행일 메타데이터)에서 기계 산출 — LLM 불개입, 가상 개체 없음."),
        ("categories", "예금/대출/카드/외환/퇴직연금/펀드/공통"),
        ("sources",
         "금융감독원 금융상품통합비교공시 Open API (202608, 은행권 19개사 215상품) · "
         "국가법령정보센터 Open API (현행 법령 7종: 은행법, 예금자보호법, 금소법, "
         "여신전문금융업법, 외국환거래법, 근로자퇴직급여 보장법, 자본시장법)"),
        ("answer_types",
         "boolean(YES/NO 균형) · numeric · span · abstain(공시 부재 — 보류 정답)"),
        ("license", "원천 데이터: 공공누리 제1유형(출처표시). 문항: 저장소 소유자."),
        ("regeneration",
         "python bench/fetch_laws.py && python bench/generate.py (결정적)"),
        ("date", "2026-08-17"),
    ]
    for r in card:
        ws.append(r)
    style(ws, {"A": 16, "B": 105})

    ws = wb.create_sheet("questions")
    ws.append(["id", "qtype", "category", "subcategory", "template", "question",
               "answer_type", "answer", "unit", "evidence_source",
               "evidence_locator", "evidence_text"])
    for r in rows:
        ws.append([r["id"], r["qtype"], r["category"], r["subcategory"],
                   r["template"], r["question"], r["answer_type"], r["answer"],
                   r.get("unit", ""), r["evidence"]["source"],
                   r["evidence"]["locator"], r["evidence"].get("text", "")])
    style(ws, {"A": 10, "B": 8, "C": 9, "D": 16, "E": 18, "F": 60, "G": 9,
               "H": 24, "I": 6, "J": 34, "K": 26, "L": 42})

    ws = wb.create_sheet("statistics")
    ws.append(["구분", "값", "수식"])
    stats = [
        ("총 문항", "=COUNTA(questions!A:A)-1"),
        ("공시형", '=COUNTIF(questions!B:B,"공시형")'),
        ("규정형", '=COUNTIF(questions!B:B,"규정형")'),
        ("boolean", '=COUNTIF(questions!G:G,"boolean")'),
        ("  YES", '=COUNTIF(questions!H:H,"YES")'),
        ("  NO", '=COUNTIF(questions!H:H,"NO")'),
        ("numeric", '=COUNTIF(questions!G:G,"numeric")'),
        ("span", '=COUNTIF(questions!G:G,"span")'),
        ("abstain", '=COUNTIF(questions!G:G,"abstain")'),
    ]
    for name, formula in stats:
        ws.append([name, formula, ""])
    ws.append([])
    ws.append(["카테고리", "문항 수", ""])
    for cat in ("예금", "대출", "카드", "외환", "퇴직연금", "펀드", "공통"):
        ws.append([cat, f'=COUNTIF(questions!C:C,"{cat}")', ""])
    style(ws, {"A": 16, "B": 12, "C": 10})

    OUTPUT.parent.mkdir(exist_ok=True)
    wb.save(OUTPUT)
    print(f"저장: {OUTPUT.relative_to(REPO_ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
