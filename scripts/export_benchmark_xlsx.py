"""Bank-Onto 벤치마크 4종 → 엑셀(HuggingFace 데이터셋 카드 스타일) 내보내기.

시트 구성:
  dataset_card   데이터셋 이름·버전·출처·방법론·라이선스
  curated        큐레이션 검색 벤치마크 18문항 + 환각유도(NEG) 2문항
  mcq            객관식 4지선다 18문항 + NEG 2문항 (결정적 채점용)
  market         시장 데이터 벤치마크 8문항 (적재 데이터에서 SPARQL로 자동 생성)
  krfinreg_bank  KR-FinReg-QA 은행 이식 76문항 (YES/NO/ABSTAIN 판정형)
  statistics     서브셋별 문항 수·유형 집계 (수식)
  baselines      현재 3-way 검색 재현율 (docs/*.md 리포트 기준)

실행: python scripts/export_benchmark_xlsx.py
출력: docs/bank-onto-benchmark-v1.xlsx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

OUTPUT = REPO_ROOT / "docs" / "bank-onto-benchmark-v1.xlsx"

HEADER_FILL = PatternFill("solid", start_color="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(vertical="top", wrap_text=True)


def style_sheet(ws, widths: dict[str, int]):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "A2"


def sheet_dataset_card(wb: Workbook):
    ws = wb.active
    ws.title = "dataset_card"
    rows = [
        ("항목", "내용"),
        ("dataset_name", "Bank-Onto-Bench"),
        ("version", "1.0"),
        ("description",
         "FIBO 스타일 한국 은행권 온톨로지 기반 3-way RAG 벤치마크 "
         "(벡터 RAG vs Graph RAG(LPG) vs 온톨로지 RAG). "
         "검색 계층(근거 재현율, LLM 없이 결정적)과 답변 계층(MCQ·판정형)을 분리 측정."),
        ("subsets",
         "curated(18+NEG 2) / mcq(18+NEG 2) / market(8, 자동생성) / "
         "krfinreg_bank(76, 판정형)"),
        ("knowledge_base",
         "ontology/*.ttl + examples/*.ttl + data/market-instances.ttl "
         "(금감원 금융상품통합비교공시 202608, 은행권 topFinGrpNo=020000, "
         "은행 19개·상품 215개·금리옵션 660개)"),
        ("gold_policy",
         "curated: 수작업 gold evidence 문자열. market: 적재 데이터에서 SPARQL로 "
         "자동 산출(데이터 갱신 시 정답 자동 갱신). krfinreg_bank: 원본(KR-FinReg-QA, "
         "202607 공시)의 근거 문구를 202608 공시에서 재검증 후 이식 — 문구가 소실된 "
         "문항은 자동 탈락(정답 부패 방지). 지식층은 정답 산출에 관여하지 않음."),
        ("scoring",
         "검색 계층: 근거 재현율(gold 문자열의 컨텍스트 포함 여부, 결정적). "
         "답변 계층: MCQ 정답 기호 일치 + KR-FinReg YES/NO/ABSTAIN 판정 일치(결정적), "
         "보조로 LLM-as-judge."),
        ("provenance",
         "curated·mcq: 본 저장소 수작업. market: finlife.fss.or.kr Open API. "
         "krfinreg_bank: KR-FinReg-QA v1(석사논문 벤치마크, 동일 API 출처)에서 "
         "은행 문항 80개 중 76개 이식(4개는 상품이 202608 공시에서 제외되어 탈락)."),
        ("determinism",
         "모든 검색 계층 평가는 PYTHONHASHSEED와 무관하게 재현됨 "
         "(tests/test_retrievers.py가 검증)."),
        ("license_note",
         "시장 데이터는 금융감독원 공공데이터(공공누리). 문항 저작은 저장소 소유자."),
        ("generated_by", "python scripts/export_benchmark_xlsx.py"),
        ("date", "2026-08-17"),
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws, {"A": 18, "B": 100})


def sheet_curated(wb: Workbook):
    from eval.benchmark_questions import NEGATIVE_QUESTIONS, QUESTIONS

    ws = wb.create_sheet("curated")
    ws.append(["id", "category", "question", "gold_evidence", "reference_answer"])
    for q in QUESTIONS:
        ws.append([q["id"], q["category"], q["question"],
                   " || ".join(q["gold"]), q["answer"]])
    for q in NEGATIVE_QUESTIONS:
        ws.append([q["id"], q["category"], q["question"],
                   "(없음 — 빈 컨텍스트가 정답)", q["answer"]])
    style_sheet(ws, {"A": 10, "B": 10, "C": 45, "D": 40, "E": 45})


def sheet_mcq(wb: Workbook):
    from eval.mcq_questions import MCQ_NEGATIVE, MCQ_QUESTIONS

    ws = wb.create_sheet("mcq")
    ws.append(["id", "question", "option_a", "option_b", "option_c", "option_d",
               "answer"])
    for q in MCQ_QUESTIONS + MCQ_NEGATIVE:
        opts = q["options"] + [""] * (4 - len(q["options"]))
        ws.append([q["id"], q["question"], *opts, "ABCD"[q["answer"]]])
    style_sheet(ws, {"A": 10, "B": 40, "C": 28, "D": 28, "E": 28, "F": 28, "G": 8})


def sheet_market(wb: Workbook):
    from eval.run_market_eval import build_market_questions
    from knowledge.kb import BankKnowledgeBase

    ws = wb.create_sheet("market")
    ws.append(["id", "category", "question", "gold_answer", "gold_sparql"])
    for q in build_market_questions(BankKnowledgeBase()):
        gold = " || ".join(str(g) for g in q["gold"])
        ws.append([q["id"], q["category"], q["question"], gold,
                   " ".join(q["sparql"].split())])
    style_sheet(ws, {"A": 12, "B": 12, "C": 40, "D": 55, "E": 60})


def sheet_krfinreg(wb: Workbook):
    ws = wb.create_sheet("krfinreg_bank")
    ws.append(["id", "group", "bank", "product_name", "product_code", "question",
               "verdict", "gold_evidence", "basis", "provenance"])
    path = REPO_ROOT / "data" / "krfinreg-bank-questions.jsonl"
    for line in path.read_text().splitlines():
        q = json.loads(line)
        ws.append([q["id"], q["group"], q["bank"], q["product_name"],
                   q["product_code"], q["question"], q["verdict"],
                   " || ".join(q["gold"]) or "(없음 — 보류가 정답)",
                   q["basis"], q["provenance"]])
    style_sheet(ws, {"A": 9, "B": 11, "C": 16, "D": 26, "E": 20, "F": 55,
                     "G": 10, "H": 40, "I": 40, "J": 30})


def sheet_statistics(wb: Workbook):
    ws = wb.create_sheet("statistics")
    ws.append(["서브셋", "문항 수", "비고"])
    ws.append(["curated", "=COUNTA(curated!A:A)-1",
               "검색 계층 (근거 재현율) — NEG 2문항 포함"])
    ws.append(["mcq", "=COUNTA(mcq!A:A)-1", "답변 계층 (결정적 채점)"])
    ws.append(["market", "=COUNTA(market!A:A)-1",
               "검색 계층 — 적재 데이터에서 자동 생성"])
    ws.append(["krfinreg_bank", "=COUNTA(krfinreg_bank!A:A)-1",
               "검색+답변 계층 (판정형)"])
    ws.append(["합계", "=SUM(B2:B5)", ""])
    ws.append([])
    ws.append(["krfinreg_bank 판정 분포", "", ""])
    for verdict in ("YES", "NO", "ABSTAIN"):
        ws.append([verdict, f'=COUNTIF(krfinreg_bank!G:G,"{verdict}")', ""])
    ws.append([])
    ws.append(["krfinreg_bank 그룹 분포", "", ""])
    for group in ("Q1_통제", "Q3_조건", "Q4_부재", "Q5_비적용"):
        ws.append([group, f'=COUNTIF(krfinreg_bank!B:B,"{group}")', ""])
    style_sheet(ws, {"A": 24, "B": 12, "C": 45})


def sheet_baselines(wb: Workbook):
    ws = wb.create_sheet("baselines")
    ws.append(["서브셋", "지표", "벡터 RAG", "Graph RAG(LPG)", "온톨로지 RAG",
               "출처"])
    rows = [
        ("curated (18문항)", "근거 재현율", 0.61, 0.89, 1.00,
         "docs/benchmark-report.md"),
        ("curated NEG (2문항)", "환각유도 컨텍스트(자/문항)", 2020, 0, 0,
         "docs/benchmark-report.md"),
        ("market (8문항)", "근거 재현율", 0.10, 0.78, 1.00,
         "docs/market-benchmark-report.md"),
        ("krfinreg_bank (판정형 71문항)", "근거 재현율", 0.88, 1.00, 1.00,
         "docs/krfinreg-benchmark-report.md"),
    ]
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0%"
    ws.append([])
    ws.append(["주: 검색 계층(LLM 없이 결정적) 결과. 답변 계층(run_answer_eval.py)은 "
               "아직 미실행 — ANTHROPIC_API_KEY 필요.", "", "", "", "", ""])
    style_sheet(ws, {"A": 28, "B": 24, "C": 12, "D": 15, "E": 14, "F": 32})


def main() -> int:
    wb = Workbook()
    sheet_dataset_card(wb)
    sheet_curated(wb)
    sheet_mcq(wb)
    sheet_market(wb)
    sheet_krfinreg(wb)
    sheet_statistics(wb)
    sheet_baselines(wb)
    OUTPUT.parent.mkdir(exist_ok=True)
    wb.save(OUTPUT)
    print(f"저장: {OUTPUT.relative_to(REPO_ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
